from __future__ import annotations

from itertools import islice
from pathlib import Path

from openpyxl import Workbook, load_workbook


SOURCE = Path("file/new_modified.xlsx")
OUTPUT_ROOT = Path("file/test_cases")
MAX_SCAN_ROWS = 20
ROWS_TO_COPY_PER_SHEET = 80

SOURCE_GROUPS = {
    "cost": {"AIR COST TRN"},
    "cash": {"CASH x SAle", "CASH X Re"},
    "spyj": {"SPYJ SALE", "SPJY Refund"},
}


def is_clean_header_row(row: tuple) -> bool:
    filled = [cell for cell in row if cell is not None and str(cell).strip()]
    return len(filled) >= 2 and len(filled) >= len(row) / 2


def detect_header_row(ws) -> int:
    for index, row in enumerate(
        ws.iter_rows(min_row=1, max_row=MAX_SCAN_ROWS, values_only=True),
        start=1,
    ):
        if is_clean_header_row(row):
            return index
    return 1


def non_empty_rows(rows) -> list[tuple]:
    return [row for row in rows if any(cell is not None and str(cell).strip() for cell in row)]


def load_source_data(source: Path) -> dict[str, dict]:
    wb = load_workbook(source, read_only=True, data_only=False)
    source_data: dict[str, dict] = {}

    for ws in wb.worksheets:
        header_row = detect_header_row(ws)
        preheader = []
        if header_row > 1:
            preheader = list(
                ws.iter_rows(
                    min_row=1,
                    max_row=header_row - 1,
                    max_col=ws.max_column,
                    values_only=True,
                )
            )

        header = next(
            ws.iter_rows(
                min_row=header_row,
                max_row=header_row,
                max_col=ws.max_column,
                values_only=True,
            )
        )
        rows = non_empty_rows(
            islice(
                ws.iter_rows(
                    min_row=header_row + 1,
                    max_col=ws.max_column,
                    values_only=True,
                ),
                ROWS_TO_COPY_PER_SHEET,
            )
        )

        source_data[ws.title] = {
            "preheader": preheader,
            "header": header,
            "rows": rows,
        }

    wb.close()
    return source_data


def create_workbook(path: Path, source_data: dict[str, dict], rows_by_sheet: dict[str, list[tuple]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, sheet in source_data.items():
        ws = wb.create_sheet(sheet_name)
        for row in sheet["preheader"]:
            ws.append(row)
        ws.append(sheet["header"])
        for row in rows_by_sheet.get(sheet_name, []):
            ws.append(row)
        ws.freeze_panes = f"A{len(sheet['preheader']) + 2}"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def rows_for_groups(source_data: dict[str, dict], groups: set[str], start: int, count: int) -> dict[str, list[tuple]]:
    allowed_sheets = set().union(*(SOURCE_GROUPS[group] for group in groups))
    return {
        sheet_name: sheet["rows"][start : start + count]
        for sheet_name, sheet in source_data.items()
        if sheet_name in allowed_sheets
    }


def build_case_1(source_data: dict[str, dict]) -> list[Path]:
    paths = []
    for part in range(2):
        rows_by_sheet = {
            sheet_name: sheet["rows"][part * 10 : part * 10 + 10]
            for sheet_name, sheet in source_data.items()
        }
        path = OUTPUT_ROOT / "case_1_row_split_same_headers" / f"case_1_workbook_{part + 1:02}.xlsx"
        create_workbook(path, source_data, rows_by_sheet)
        paths.append(path)
    return paths


def build_case_2(source_data: dict[str, dict]) -> list[Path]:
    scenarios = [
        ("cost_rows", {"cost"}, 20),
        ("cash_rows", {"cash"}, 20),
        ("spyj_rows", {"spyj"}, 20),
    ]
    paths = []
    for name, groups, start in scenarios:
        rows_by_sheet = rows_for_groups(source_data, groups, start=start, count=12)
        path = OUTPUT_ROOT / "case_2_source_groups_split" / f"case_2_{name}.xlsx"
        create_workbook(path, source_data, rows_by_sheet)
        paths.append(path)
    return paths


def build_case_3(source_data: dict[str, dict]) -> list[Path]:
    scenarios = [
        ("cost_cash_no_spyj", {"cost", "cash"}, 32),
        ("cost_spyj_no_cash", {"cost", "spyj"}, 44),
        ("cost_only", {"cost"}, 56),
        ("cash_spyj_no_cost", {"cash", "spyj"}, 68),
    ]
    paths = []
    for name, groups, start in scenarios:
        rows_by_sheet = rows_for_groups(source_data, groups, start=start, count=12)
        path = OUTPUT_ROOT / "case_3_source_presence_filters" / f"case_3_{name}.xlsx"
        create_workbook(path, source_data, rows_by_sheet)
        paths.append(path)
    return paths


def write_manifest(paths: list[Path]) -> None:
    lines = [
        "# Multi-workbook reconciliation test cases",
        "",
        f"Source workbook: `{SOURCE}`",
        "All fixture rows are copied from the source workbook. Sheet names, title rows, and headers are preserved.",
        "",
        "## Case 1: row split with same headers",
        "Upload both workbooks together. Every workbook contains all five source sheets with 10 real data rows per sheet.",
        "",
        "## Case 2: source groups split across workbooks",
        "Upload all three workbooks together. Cost, CASH X, and SPYJ rows are placed in separate workbooks to confirm same-sheet data is merged across files.",
        "",
        "## Case 3: source presence filters",
        "Upload one workbook at a time, or upload the folder together, to test missing-source filters. Each populated sheet has 12 real data rows.",
        "",
        "## Files",
    ]
    lines.extend(f"- `{path}`" for path in paths)
    (OUTPUT_ROOT / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source workbook not found: {SOURCE}")

    source_data = load_source_data(SOURCE)
    paths = []
    paths.extend(build_case_1(source_data))
    paths.extend(build_case_2(source_data))
    paths.extend(build_case_3(source_data))
    write_manifest(paths)

    print(f"Created {len(paths)} source-derived workbooks under {OUTPUT_ROOT}")
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
